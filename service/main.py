import os
import io
import json
import numpy as np
from fastapi import FastAPI, HTTPException, UploadFile, File
from fastapi.responses import StreamingResponse
from docx import Document
from openai import OpenAI

# Configuración
OPENAI_API_KEY = os.getenv("OPENAI_API_KEY")
EMBEDDINGS_FILE = "embeddings.json"  # Debe estar en la raíz de periciales-review
MODEL_EMBED = "text-embedding-ada-002"
MODEL_CHAT = "gpt-4o"

client = OpenAI(api_key=OPENAI_API_KEY)
app = FastAPI(title="Informe Review Agent (In-Memory RAG Completo)")

# Carga embeddings y textos completos al iniciar
def load_data(path):
    with open(path, 'r', encoding='utf-8') as f:
        data = json.load(f)
    embeddings = np.array([item['embedding'] for item in data])
    texts = [item['text'] for item in data]
    norms = np.linalg.norm(embeddings, axis=1)
    return embeddings, texts, norms

embeddings, texts, norms = load_data(os.path.join(os.path.dirname(__file__), "..", EMBEDDINGS_FILE))

# Función de similitud coseno
def top_k_similar(query_emb, k=5):
    qnorm = np.linalg.norm(query_emb)
    sims = embeddings.dot(query_emb) / ((norms * qnorm) + 1e-10)
    idxs = np.argsort(-sims)[:k]
    return [texts[i] for i in idxs]

# Extraer texto de DOCX
def extract_text_from_docx(bytes_data):
    doc = Document(io.BytesIO(bytes_data))
    parts = [p.text for p in doc.paragraphs if p.text]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    parts.append(cell.text)
    return "\n".join(parts)

# Rellenar plantilla DOCX
def fill_docx_fields(bytes_data, fields):
    doc = Document(io.BytesIO(bytes_data))
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    for run in para.runs:
                        for key, val in fields.items():
                            if key in run.text:
                                run.text = run.text.replace(key, val)
    out = io.BytesIO()
    doc.save(out)
    out.seek(0)
    return out

@app.post("/review")
async def review(report_file: UploadFile = File(...), top_k: int = 5):
    # 1) Leer informe
    file_bytes = await report_file.read()
    raw_text = extract_text_from_docx(file_bytes)

    # 2) Embedding del nuevo informe
    emb_resp = client.embeddings.create(input=raw_text, model=MODEL_EMBED)
    query_emb = np.array(emb_resp.data[0].embedding)

    # 3) Recuperar textos similares
    retrieved_texts = top_k_similar(query_emb, k=top_k)

    # 4) Llamar a GPT para extracción y corrección
    system_msg = (
        "Eres un asistente experto en informes periciales. "
        "Usa estos ejemplos completos para guiar tu corrección:"
    )
    messages = [
        {"role": "system", "content": system_msg},
        {"role": "system", "name": "context_reports", "content": json.dumps(retrieved_texts, ensure_ascii=False)},
        {"role": "user", "content": raw_text}
    ]
    resp = client.chat.completions.create(model=MODEL_CHAT, messages=messages)
    try:
        result = json.loads(resp.choices[0].message.content)
    except Exception:
        raise HTTPException(status_code=500, detail="Respuesta de GPT no válida JSON")
    fields = result.get("fields", {})

    # 5) Rellenar y devolver el DOCX corregido
    doc_out = fill_docx_fields(file_bytes, fields)
    filename = report_file.filename.replace('.docx', '_revisado.docx')
    return StreamingResponse(
        doc_out,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )

if __name__ == "__main__":
    import uvicorn
    port = int(os.getenv("PORT", 8080))
    uvicorn.run(app, host="0.0.0.0", port=port)
